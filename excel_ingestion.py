"""Bounded XLSX workbook parsing, sheet selection, column inspection and mapping."""
import io
import json

from analysis_service import AnalysisError
from config import MAX_EXCEL_BYTES, MAX_EXCEL_COLUMNS, MAX_BATCH_RESPONSES

ALIASES = {
    "text_column": {"feedback", "comment", "review", "response", "text"},
    "date_column": {"date", "timestamp", "submitted_at"},
    "category_column": {"category", "department", "service", "type"},
}


def parse_excel(stream, sheet_name=None):
    """Parse an XLSX workbook and return (sheet_names, columns, rows).

    Only openpyxl is used; macros are not executed. When sheet_name is
    provided the matching sheet is parsed; otherwise the first worksheet is.
    """
    try:
        import openpyxl
    except ImportError:
        raise AnalysisError("Excel support requires the openpyxl package.")

    payload = stream.read(MAX_EXCEL_BYTES + 1)
    if len(payload) > MAX_EXCEL_BYTES:
        raise AnalysisError(f"Excel file exceeds {MAX_EXCEL_BYTES} bytes.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise AnalysisError("Could not read the Excel workbook. Ensure it is a valid .xlsx file.") from exc

    try:
        sheet_names = wb.sheetnames
        if not sheet_names:
            raise AnalysisError("Excel workbook contains no sheets.")

        if sheet_name and sheet_name not in sheet_names:
            raise AnalysisError(f"Sheet '{sheet_name}' was not found in the workbook.")
        ws = wb[sheet_name if sheet_name else sheet_names[0]]
        all_rows = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(row)
        if not all_rows:
            raise AnalysisError("The selected sheet contains no data.")

        header = all_rows[0]
        columns = []
        for i, cell in enumerate(header):
            if cell is None:
                columns.append(f"Column {i + 1}")
            else:
                columns.append(str(cell).strip())

        if not any(c for c in columns):
            raise AnalysisError("The selected sheet has no usable column headers.")

        if len(columns) > MAX_EXCEL_COLUMNS:
            raise AnalysisError(f"Excel sheet has {len(columns)} columns; at most {MAX_EXCEL_COLUMNS} are supported.")

        if len({c.casefold() for c in columns if c}) != len([c for c in columns if c]):
            raise AnalysisError("Excel sheet requires unique, non-empty column names.")

        rows = []
        for row_num, row in enumerate(all_rows[1:], start=1):
            if len(rows) >= MAX_BATCH_RESPONSES:
                raise AnalysisError(f"Excel sheet exceeds {MAX_BATCH_RESPONSES} data records.")
            row_dict = {}
            for i, col_name in enumerate(columns):
                cell_value = row[i] if i < len(row) else None
                if cell_value is None:
                    row_dict[col_name] = ""
                else:
                    row_dict[col_name] = str(cell_value)
            rows.append(row_dict)

        if not rows:
            raise AnalysisError("The selected sheet has a header but no data records.")

        return sheet_names, columns, rows
    finally:
        wb.close()


def inspect_excel_columns(sheet_names, columns, rows):
    """Return column inspection results for Excel, including available sheets."""
    candidates = {role: [name for name in columns if name.casefold() in aliases]
                  for role, aliases in ALIASES.items()}
    suggested = {role: values[0] if len(values) == 1 else None for role, values in candidates.items()}
    return {"columns": columns, "row_count": len(rows), "suggested_mapping": suggested,
            "candidates": candidates, "requires_selection": suggested["text_column"] is None,
            "sheets": sheet_names}


def map_excel(columns, rows, options):
    """Map Excel columns to response records using the same logic as PDF."""
    inspection = inspect_excel_columns([], columns, rows)
    selected = {}
    for role in (*ALIASES, "id_column", "source_column"):
        name = options.get(role, inspection["suggested_mapping"].get(role))
        if name and name not in columns:
            raise AnalysisError(f"Selected {role} is not a column.", inspection)
        selected[role] = name or None
    if not selected["text_column"]:
        raise AnalysisError("Select the feedback column using text_column.", inspection)
    role_columns = [name for name in selected.values() if name]
    if len(role_columns) != len(set(role_columns)):
        raise AnalysisError("Select a different column for each mapped field.")
    try:
        metadata_columns = json.loads(options.get("metadata_columns", "[]"))
    except (ValueError, TypeError, RecursionError) as exc:
        raise AnalysisError("metadata_columns must be a JSON array of column names.") from exc
    if (not isinstance(metadata_columns, list) or any(not isinstance(name, str) or name not in columns for name in metadata_columns)
            or len(metadata_columns) != len(set(metadata_columns))):
        raise AnalysisError("metadata_columns must contain unique existing column names.")
    normalized = []
    for row in rows:
        entry = {"text": row[selected["text_column"]]}
        for role in ("id", "date", "category", "source"):
            column = selected[role + "_column"]
            if column and row[column].strip():
                entry[role] = row[column]
        if metadata_columns:
            entry["metadata"] = {name: row[name] for name in metadata_columns}
        normalized.append(entry)
    return normalized
